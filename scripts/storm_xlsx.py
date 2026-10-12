#!/usr/bin/env python3
"""Spreadsheet of the June-July storm history by marine forecast area (data/storm_history), for reading outside the planner.

Usage: storm_xlsx.py OUT.xlsx
Sheets: Summary (one row per area: how often, how likely in a 16-day crossing, how long), Duration (length spread and
how many blows of each length), Timing (by half month and by decade), By year (each area and year), Notes (sources and method).
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
D = ROOT / "data" / "storm_history"
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E79")
WRAP = Alignment(wrap_text=True, vertical="top")
LEVELS = (("sc", "Small craft+"), ("gale", "Gale+"), ("storm", "Storm"))


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(10, min(40, len(str(h)) + 2)))
    ws.row_dimensions[1].height = 45
    return ws


def source(a):
    o = a.get("observed")
    if a["best_source"] == "observed" and o:
        return ", ".join(s.title() for s in o["stations"])
    return a["best_source"]


def main():
    out = Path(sys.argv[1])
    S = json.loads((D / "summary.json").read_text())
    A = S["areas"]
    wb = Workbook()
    wb.remove(wb.active)

    hdr = ["Area", "Record used", "Years", "Small craft+ blows per Jun-Jul", "Chance of a small craft+ blow in 16 days",
           "Days with small craft+ per Jun-Jul", "Share of hours at small craft+ (%)", "Years with any small craft+ (%)",
           "Typical small craft+ length (h)", "1 in 10 lasts over (h)", "Longest (h)",
           "Gale+ blows per Jun-Jul", "Chance of a gale+ in 16 days", "Gale+ typical length (h)", "Storm blows per Jun-Jul",
           "ERA5 1976-2025: small craft+ per Jun-Jul", "ERA5 1976-2025: gale+ per Jun-Jul", "Note"]
    rows = []
    for a in A:
        b, e = a["best"], a.get("corrected") or a["era5_raw"]
        sc, g, st = b["sc"], b["gale"], b["storm"]
        rows.append([a["area"], source(a), a["best_years"], sc["events_per_jj"], round(sc["p_in_16_days"] * 100), sc["days_per_jj"],
                     sc["share_of_hours_pct"] if a["best_source"] != "NWS warnings issued" else None, sc["years_with_any_pct"],
                     sc["dur_h"]["median"], sc["dur_h"]["p90"], sc["dur_h"]["max"], g["events_per_jj"], round(g["p_in_16_days"] * 100),
                     g["dur_h"]["median"], st["events_per_jj"], e["sc"]["events_per_jj"], e["gale"]["events_per_jj"], a.get("note", "")])
    ws = sheet(wb, "Summary", hdr, rows, {1: 34, 2: 30, 3: 20, 18: 40})
    for r in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in r:
            c.number_format = '0"%"'
    for r in ws.iter_rows(min_row=2, min_col=13, max_col=13):
        for c in r:
            c.number_format = '0"%"'

    hdr = ["Area", "Level", "Blows per Jun-Jul", "25% last under (h)", "Median (h)", "75% last under (h)", "90% last under (h)",
           "Longest (h)", "Number of blows (all years) lasting under 12 h", "12-24 h", "1-2 days", "2-3 days", "3+ days"]
    rows = []
    for a in A:
        for k, lab in LEVELS:
            s = a["best"][k]
            d, L = s["dur_h"], s["longest_by_days"]
            rows.append([a["area"], lab, s["events_per_jj"], d["p25"], d["median"], d["p75"], d["p90"], d["max"], *L.values()])
    sheet(wb, "Duration", hdr, rows, {1: 34})

    halves = list(A[0]["best"]["sc"]["by_half_month_share_pct"])
    decs = list(A[0]["best"]["sc"]["by_decade_events"])
    hdr = ["Area", "Level"] + [f"% of hours at level, {h}" for h in halves] + [f"Blows per Jun-Jul, {d}" for d in decs]
    rows = []
    for a in A:
        for k, lab in LEVELS[:2]:
            s = a["best"][k]
            rows.append([a["area"], lab, *s["by_half_month_share_pct"].values(),
                         *[v if v else None for v in s["by_decade_events"].values()]])
    sheet(wb, "Timing", hdr, rows, {1: 34})

    with open(D / "area_year.csv") as f:
        cr = list(csv.DictReader(f))
    names = {"best_source": "Record used that year", "best_sc_events": "Small craft+ blows", "best_gale_events": "Gale+ blows",
             "best_storm_events": "Storm blows", "best_sc_hours": "Hours at small craft+", "best_longest_h": "Longest blow (h)",
             "best_max_kn": "Highest wind (kn)", "era5_sc_events": "ERA5 small craft+ blows", "era5_gale_events": "ERA5 gale+ blows",
             "era5_storm_events": "ERA5 storm blows", "era5_sc_hours": "ERA5 hours at small craft+", "era5_longest_h": "ERA5 longest (h)",
             "era5_max_kn": "ERA5 highest wind (kn)"}
    cols = ["area", "year"] + list(names)
    num = lambda v: (float(v) if "." in v else int(v)) if v not in ("", None) else None
    rows = [[r["area"], int(r["year"]), r.get("best_source", "")] + [num(r.get(c, "")) for c in cols[3:]] for r in cr]
    sheet(wb, "By year", ["Area", "Year"] + list(names.values()), rows, {1: 34, 3: 28})

    notes = [
        ["Period", "June 1 to July 31 each year, 1976-2025, times Pacific daylight time."],
        ["Areas", "The Environment Canada marine forecast areas the route passes through, Juan de Fuca Strait east entrance to Dixon "
                  "Entrance east, plus Hecate Strait (outside route only), the inside channels from Finlayson to Grenville (no forecast "
                  "area of its own) and US zone PKZ036 Clarence Strait for the run to Ketchikan."],
        ["Thresholds", "Small craft 20 kn sustained in Canada, 23 kn in Alaska (the levels the warnings are issued at); gale 34 kn; storm 48 kn."],
        ["Blow", "A run of hours at small craft strength or more. Gaps up to 6 h are merged into one blow and blows under 2 h are "
                 "dropped. Its level is the highest it reached. Length runs from first to last hour at the level."],
        ["Chance in 16 days", "Share of 16-day windows inside June-July (about the length of a crossing) that contain at least one "
                              "blow starting in them."],
        ["Record used", "Each area uses the best record it has: hourly wind from Environment Canada lightstations and automatic "
                        "stations in the area (NDBC stations in Alaska), counting years at least 40% complete after carrying each "
                        "report forward up to 3 h; else the area's MEDS weather buoy; else, in Clarence Strait, the warnings the US "
                        "National Weather Service actually issued since 2005; else ERA5."],
        ["Why not 50 years everywhere", "Hourly station wind is mostly archived from the mid 1990s, so the observed records cover "
                                        "about 25-35 years. ERA5 reanalysis covers all 50 years and is shown alongside, but it is far "
                                        "too smooth inside the channels (it finds almost no small craft winds in Johnstone Strait or "
                                        "the Central Coast), so it is only a cross-check there."],
        ["Exposed stations", "Lightstations sit on exposed points and read higher than the middle of a sheltered channel; treat the "
                             "numbers as what the forecast area sees, not every bay in it."],
        ["Source note", S.get("note", "")],
        ["Generated", S.get("generated", "")],
    ]
    ws = sheet(wb, "Notes", ["Item", "Detail"], notes, {1: 26, 2: 120})
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = WRAP
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
